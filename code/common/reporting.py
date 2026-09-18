"""Figure and table furniture shared by the reporting stages.
"""

from __future__ import annotations

import sys
from pathlib import Path

import numpy as np
import pandas as pd
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt

sys.path.insert(0, str(Path(__file__).resolve().parent))
import survival as sv

# ── Reported strata ──────────────────────────────────────────────────────────

STRATA = [
    ("education", ["1", "2", "3", "4"]),
    ("nssec_class8", [str(x) for x in range(1, 9)]),
    ("tenure", ["1", "2", "3", "4", "5", "6"]),
    ("household_size", ["1", "2"]),
    ("ruralurban", ["1", "2", "3", "4"]),
    ("imd_decile", [str(x) for x in range(1, 11)]),
]

STRATUM_LEVEL_LABELS = {
    "education": {"1": "Level 1", "2": "Level 2", "3": "Level 3",
                  "4": "Level 4"},
    "nssec_class8": {"1": "Higher managerial and professional",
                     "2": "Lower managerial and professional",
                     "3": "Intermediate",
                     "4": "Small employers and own account",
                     "5": "Lower supervisory and technical",
                     "6": "Semi-routine",
                     "7": "Routine",
                     "8": "Never worked, long-term unemployed"},
    "tenure": {"1": "Owned outright", "2": "Owned with mortgage/loan",
               "3": "Shared ownership", "4": "Social rented",
               "5": "Private rented", "6": "Living rent free", "7": "Other"},
    "household_size": {"1": "1", "2": "2+"},
    "ruralurban": {"1": "Urban", "2": "Town and Fringe", "3": "Village",
                   "4": "Hamlet/Isolated Dwelling"},
    "sex": {"1": "Male", "2": "Female"},
    "ethnicity": {"1": "White", "2": "Mixed", "3": "Asian", "4": "Black",
                  "5": "Chinese", "6": "Other"},
    "econstatus": {"1": "Employed", "2": "Unemployed", "3": "Retired",
                   "4": "Other"},
    "health": {"1": "Excellent", "2": "Good", "3": "Fair", "4": "Poor"},
}

STRATA_LABELS = {
    "education": "Education",
    "nssec_class8": "Socioeconomic class",
    "tenure": "Housing tenure",
    "household_size": "Household size",
    "ruralurban": "Urbanisation",
    "imd_decile": "Neighbourhood deprivation",
}

LEVEL_ANNOTATIONS = {
    "imd_decile": {"1": "1 (most deprived)", "10": "10 (least deprived)"},
    "education": {"1": "Level 1 (lowest)", "4": "Level 4 (highest)"},
}

PARITY_STRATA = ["education", "tenure", "ruralurban", "imd_decile"]

FALSE_NEGATIVE_RATE_LEVELS = {
    "education": ["1", "4"],
    "nssec_class8": ["1", "8"],
    "tenure": ["3", "4"],
    "household_size": ["1", "2"],
    "ruralurban": ["2", "4"],
    "imd_decile": ["1", "10"],
}

_declared = {v for v, _ in STRATA}
for _name, _keys in (("STRATA_LABELS", set(STRATA_LABELS)),
                     ("PARITY_STRATA", set(PARITY_STRATA)),
                     ("FALSE_NEGATIVE_RATE_LEVELS",
                      set(FALSE_NEGATIVE_RATE_LEVELS))):
    if not _keys <= _declared:
        raise RuntimeError(
            f"{_name} names " + ", ".join(sorted(_keys - _declared))
            + ", which STRATA does not declare.")

for _var, _levels in STRATA:
    _modelled = sv.CATEGORY_LEVEL_ORDERS.get(_var)
    if _modelled is not None and not set(_levels) <= set(_modelled):
        raise RuntimeError(
            f"STRATA gives {_var} the levels "
            + ", ".join(sorted(set(_levels) - set(_modelled)))
            + ", which the harmonisation does not produce. It carries "
            + ", ".join(sorted(_modelled)) + ".")

def level_label(var, level):
    annotated = LEVEL_ANNOTATIONS.get(var, {}).get(str(level))
    if annotated is not None:
        return annotated
    return STRATUM_LEVEL_LABELS.get(str(var), {}).get(str(level), str(level))

def as_level(series):
    if pd.api.types.is_float_dtype(series):
        return series.map(lambda v: "" if pd.isna(v) else f"{int(round(v))}")
    return series.astype(str)

def stratum_rows(df):
    rows = []
    for var, levels in STRATA:
        rows.append({"kind": "group", "var": var, "level": "",
                     "label": STRATA_LABELS.get(var, var)})
        vals = as_level(df[var])
        for level in levels:
            if not (vals == level).any():
                continue
            rows.append({"kind": "level", "var": var, "level": level,
                         "label": level_label(var, level)})
    return pd.DataFrame(rows)


def level_codes(df):
    out = []
    for var, levels in STRATA:
        vals = as_level(df[var]).to_numpy()
        keep = [lv for lv in levels if (vals == str(lv)).any()]
        if len(keep) < 2:
            continue
        codes = np.full(len(vals), -1, dtype=np.int64)
        for i, lv in enumerate(keep):
            codes[vals == str(lv)] = i
        out.append({"var": var, "title": STRATA_LABELS.get(var, var),
                    "levels": keep, "codes": codes})
    return out


def accumulate(codes, k, weights, mask=None):
    m = codes >= 0 if mask is None else (codes >= 0) & mask
    return np.bincount(codes[m], weights=weights[m], minlength=k)


# ── Colours and markers ──────────────────────────────────────────────────────

C_PMR = "#888888"
C_UKB = "#640404"
C_WEIGHTED = "#2C5D8A"
M_PMR, M_UKB, M_WEIGHTED = "D", "s", "o"

# ── Saving ───────────────────────────────────────────────────────────────────

def save_fig(fig, output_dir, name):
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    fig.savefig(output_dir / f"{name}.png", dpi=300, bbox_inches="tight",
                facecolor="white")
    plt.close(fig)
    print(f"  wrote {output_dir / (name + '.png')}")

def write_workbook(path, sheets):
    from openpyxl.styles import Alignment, Font

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(path, engine="openpyxl") as xl:
        for name, frame in sheets.items():
            frame.to_excel(xl, sheet_name=name, index=False)
            ws = xl.sheets[name]
            ws.freeze_panes = "A2"
            for j, col in enumerate(frame.columns, start=1):
                width = max([len(str(col))]
                            + [len(str(v)) for v in frame[col].tolist()]) + 3
                ws.column_dimensions[
                    ws.cell(row=1, column=j).column_letter].width = width
                for i in range(2, len(frame) + 2):
                    ws.cell(row=i, column=j).alignment = Alignment(
                        horizontal="left" if j <= 3 else "right")
            for j in range(1, len(frame.columns) + 1):
                ws.cell(row=1, column=j).font = Font(bold=True)
                ws.cell(row=1, column=j).alignment = Alignment(
                    horizontal="center", vertical="bottom", wrap_text=True)
    print(f"  wrote {path}")


# ── The lollipop panels ──────────────────────────────────────────────────────
def lollipop(ax, rows, series, xlabel):
    import matplotlib.ticker as mticker

    n = len(rows)
    y = np.arange(n)
    for i, row in rows.reset_index(drop=True).iterrows():
        if row["kind"] != "level":
            continue
        for col, colour, marker in series:
            v = row.get(col, np.nan)
            if not np.isfinite(v):
                continue
            ax.plot([0, v], [i, i], color=colour, linewidth=1.6, zorder=2)
            ax.scatter(v, i, c=colour, marker=marker, s=42, zorder=4,
                       edgecolors="none")

    ax.set_yticks(y)
    ax.set_yticklabels(["" if r["kind"] == "group" else "    " + r["label"]
                        for _, r in rows.reset_index(drop=True).iterrows()],
                       fontsize=8.5)
    for tick, (_, r) in zip(ax.get_yticklabels(),
                            rows.reset_index(drop=True).iterrows()):
        if r["kind"] == "group":
            tick.set_text("")
    for i, (_, r) in enumerate(rows.reset_index(drop=True).iterrows()):
        if r["kind"] == "group":
            ax.text(0.0, i, r["label"], transform=ax.get_yaxis_transform(),
                    ha="right", va="center", fontsize=9, fontweight="bold")

    ax.invert_yaxis()
    ax.set_ylim(n - 0.4, -0.8)
    ax.axvline(0, color="black", linewidth=1.6, zorder=3)
    ax.spines[["top", "right", "left"]].set_visible(False)
    ax.spines["bottom"].set_linewidth(1.6)
    ax.tick_params(axis="y", length=0)
    ax.tick_params(axis="x", width=1.6, labelsize=10)
    ax.xaxis.set_major_formatter(
        mticker.FuncFormatter(lambda v, _: f"{int(v):,}"))
    ax.set_xlabel(xlabel, fontsize=11)

SWEEP_AXIS_MAX = {5: 50.0, 10: 100.0}

def sweep_panel(ax, title, years, ylabel=None, ynumbers=True, legend=True):
    ax.axhline(0, color="#bbbbbb", linewidth=0.8)
    ax.set_xlim(0, SWEEP_AXIS_MAX[int(years)])
    ax.set_title(title, fontsize=8.5)
    ax.set_xlabel("Risk threshold (%)", fontsize=7.5)
    if ylabel:
        ax.set_ylabel(ylabel, fontsize=8)
    ax.tick_params(labelsize=7, labelleft=ynumbers)
    if legend:
        labels = ax.get_legend_handles_labels()[1]
        longest = max((len(l) for l in labels), default=0)
        ax.legend(fontsize=5 if longest > 24 else 6, frameon=False,
                  loc="upper right", handlelength=1.4, labelspacing=0.3,
                  borderaxespad=0.4)

def pad_axis(ax, values, scale=1.15):
    v = np.asarray([x for x in values if np.isfinite(x)], dtype=float)
    lo = min(float(v.min()), 0.0) * scale if len(v) else -1.0
    hi = max(float(v.max()), 0.0) * scale if len(v) else 1.0
    span = hi - lo
    ax.set_xlim(lo - 0.02 * span if lo < 0 else -0.02 * span,
                hi + 0.02 * span if hi > 0 else 0.02 * span)

def bar_with_interval(ax, y, value, lo, hi, colour, height=0.28, cap=0.07,
                      linewidth=1.2):
    if not np.isfinite(value):
        return
    ax.barh(y, value, height=height, color=colour, linewidth=0, zorder=4, left=0)
    vx0, vx1 = ax.get_xlim()
    if value < vx0:
        ax.plot([vx0], [y], marker="<", markersize=4.0, color=colour,
                clip_on=False, zorder=9)
    elif value > vx1:
        ax.plot([vx1], [y], marker=">", markersize=4.0, color=colour,
                clip_on=False, zorder=9)
    if not (np.isfinite(lo) and np.isfinite(hi)):
        return

    bar_lo, bar_hi = min(0.0, value), max(0.0, value)

    def inside(x):
        return bar_lo <= x <= bar_hi

    x0, x1 = ax.get_xlim()
    draw_lo, draw_hi = max(lo, x0), min(hi, x1)
    if draw_hi > draw_lo:
        cuts = sorted({draw_lo, draw_hi}
                      | {c for c in (bar_lo, bar_hi) if draw_lo < c < draw_hi})
        for a, b in zip(cuts, cuts[1:]):
            ax.plot([a, b], [y, y],
                    color="white" if inside((a + b) / 2.0) else colour,
                    linewidth=linewidth, zorder=7, solid_capstyle="butt")
    for x, past, marker in ((lo, lo < x0, "<"), (hi, hi > x1, ">")):
        if past:
            edge = x0 if marker == "<" else x1
            ax.plot([edge], [y], marker=marker, markersize=3.2,
                    color="white" if inside(edge) else colour,
                    clip_on=False, zorder=9)
        else:
            ax.plot([x, x], [y - cap, y + cap],
                    color="white" if inside(x) else colour,
                    linewidth=linewidth, zorder=8, solid_capstyle="butt")

def box_spine(ax, linewidth=1.0):
    for spine in ax.spines.values():
        spine.set_visible(True)
        spine.set_linewidth(linewidth)
        spine.set_color("black")
